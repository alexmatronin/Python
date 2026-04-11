import csv
from datetime import datetime
from openpyxl import Workbook

def calculate_age(birth_date_str):
   
    birth_date = datetime.strptime(birth_date_str, "%Y-%m-%d").date()
    today = datetime.today().date()
  
    age = today.year - birth_date.year - ((today.month, today.day) < (birth_date.month, birth_date.day))
    return age

def main():
    csv_filename = 'employees.csv'
    xlsx_filename = 'employees.xlsx'
    employees_data = []

    try:
        with open(csv_filename, 'r', encoding='utf-8') as file:
            reader = csv.DictReader(file)
            for row in reader:
                employees_data.append(row)
    except Exception as e:
        print("Повідомлення про відсутність, або проблеми при відкритті файлу CSV, або Ok.")
        return

    try:
        workbook = Workbook()
        
        sheet_all = workbook.active
        sheet_all.title = "all"

        sheet_younger_18 = workbook.create_sheet("younger_18")
        sheet_18_45 = workbook.create_sheet("18-45")
        sheet_45_70 = workbook.create_sheet("45-70")
        sheet_older_70 = workbook.create_sheet("older_70")
        
        if employees_data:
            headers_all = list(employees_data[0].keys())

            for col_num, header in enumerate(headers_all, 1):
                sheet_all.cell(row=1, column=col_num, value=header)
 
            for row_num, row_dict in enumerate(employees_data, 2):
                for col_num, header in enumerate(headers_all, 1):
                    sheet_all.cell(row=row_num, column=col_num, value=row_dict[header])

        headers_age = ['№', 'Прізвище', "Ім'я", 'По батькові', 'Дата народження', 'Вік']
        sheets_dict = {
            "younger_18": {"sheet": sheet_younger_18, "current_row": 2},
            "18-45": {"sheet": sheet_18_45, "current_row": 2},
            "45-70": {"sheet": sheet_45_70, "current_row": 2},
            "older_70": {"sheet": sheet_older_70, "current_row": 2}
        }
        
        for sheet_info in sheets_dict.values():
            ws = sheet_info["sheet"]
            for col_num, header in enumerate(headers_age, 1):
                ws.cell(row=1, column=col_num, value=header)

        for row_dict in employees_data:
            age = calculate_age(row_dict['Дата народження'])

            if age < 18:
                category = "younger_18"
            elif 18 <= age <= 45:
                category = "18-45"
            elif 45 < age <= 70:
                category = "45-70"
            else:
                category = "older_70"

            ws = sheets_dict[category]["sheet"]
            current_row = sheets_dict[category]["current_row"]

            row_values = [
                current_row - 1, 
                row_dict['Прізвище'],
                row_dict["Ім'я"],
                row_dict['По батькові'],
                row_dict['Дата народження'],
                age
            ]
            
            for col_num, cell_value in enumerate(row_values, 1):
                ws.cell(row=current_row, column=col_num, value=cell_value)

            sheets_dict[category]["current_row"] += 1

        workbook.save(xlsx_filename)
        print("Ok")

    except Exception as e:
        print("Повідомлення про неможливість створення XLSX файлу.")

if __name__ == "__main__":
    main()